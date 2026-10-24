# SPDX-FileCopyrightText: 2026 Malte Dreyer
# SPDX-License-Identifier: MIT
"""Reading uploaded documents and writing exports.

Uploads are read once and deleted by the caller immediately afterwards
(the file server of the UI framework would otherwise serve them to anyone
who knows the URL). Only the extracted text is kept, in session memory.
"""

from __future__ import annotations

import uuid
from pathlib import Path

READABLE = (".txt", ".md", ".docx")
EXPORT_FORMATS = ("txt", "md", "docx")


class DocumentError(ValueError):
    def __init__(self, code: str, **params: object) -> None:
        super().__init__(code)
        self.code = code
        self.params = params


def read_text(path: str | Path, max_chars: int) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in READABLE:
        raise DocumentError("doc_unsupported", suffix=suffix or "?")
    try:
        if suffix == ".docx":
            from docx import Document  # imported lazily: only needed for Word files

            text = "\n\n".join(p.text for p in Document(str(path)).paragraphs if p.text.strip())
        else:
            raw = path.read_bytes()
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("cp1252", errors="replace")
    except DocumentError:
        raise
    except Exception:  # broken or foreign file: report, never show a traceback
        raise DocumentError("doc_unreadable", name=path.name) from None
    text = text.replace("\r\n", "\n").strip()
    if not text:
        raise DocumentError("doc_empty", name=path.name)
    if len(text) > max_chars:
        raise DocumentError("doc_too_long", limit=max_chars)
    return text


def write_export(
    folder: Path,
    fmt: str,
    output: str,
    original: str | None = None,
    headings: tuple[str, str] = ("Original", "Result"),
) -> Path:
    """Write a result file with a unique name into a per-session folder."""
    if fmt not in EXPORT_FORMATS:
        raise DocumentError("doc_unsupported", suffix=fmt)
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / uuid.uuid4().hex / f"text-style-result.{fmt}"
    target.parent.mkdir()
    if fmt == "docx":
        from docx import Document

        doc = Document()
        if original is not None:
            doc.add_heading(headings[0], level=1)
            for para in original.split("\n\n"):
                doc.add_paragraph(para)
            doc.add_heading(headings[1], level=1)
        for para in output.split("\n\n"):
            doc.add_paragraph(para)
        doc.save(str(target))
        return target
    if original is not None:
        if fmt == "md":
            body = f"## {headings[0]}\n\n{original}\n\n## {headings[1]}\n\n{output}\n"
        else:
            body = f"{headings[0]}\n\n{original}\n\n{headings[1]}\n\n{output}\n"
    else:
        body = output + "\n"
    target.write_text(body, encoding="utf-8")
    return target
